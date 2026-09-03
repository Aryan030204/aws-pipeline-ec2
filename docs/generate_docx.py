"""
One-off script: builds docs/Pipeline-Documentation.docx from the same
content as README.md. Not part of the pipeline itself - re-run manually
if the documentation content changes.
"""

import os

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "images")
OUT_PATH = os.path.join(HERE, "Pipeline-Documentation.docx")

NAVY = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x47, 0x55, 0x69)


def set_heading_style(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY


def add_title(doc, text, subtitle=None):
    h = doc.add_heading(text, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].italic = True
        p.runs[0].font.color.rgb = GRAY


def add_h1(doc, text):
    doc.add_heading(text, level=1)


def add_h2(doc, text):
    doc.add_heading(text, level=2)


def add_body(doc, text):
    doc.add_paragraph(text)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_image(doc, filename, width_in=6.3):
    doc.add_picture(os.path.join(IMG_DIR, filename), width=Inches(width_in))


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph("")


def build():
    doc = Document()
    set_heading_style(doc)

    add_title(
        doc,
        "Shopify Pipeline",
        "A multi-brand Shopify data pipeline — orders, sessions, checkout events, "
        "page speed, and dashboard summaries, kept up to date automatically.",
    )

    add_h1(doc, "What it does")
    add_body(
        doc,
        "For every brand it manages, the pipeline pulls in new data, keeps that "
        "brand's own MySQL database up to date, and rebuilds dashboard tables — "
        "only for the parts that actually changed. Specifically, it:",
    )
    add_bullets(
        doc,
        [
            "Fetches new and updated Shopify orders into that brand's own MySQL database.",
            "Pulls session and traffic data via ShopifyQL (Shopify's analytics query language).",
            "Syncs checkout-initiated and buy-now events from MongoDB.",
            "Syncs page-speed (“web vitals”) data from MongoDB.",
            "Rebuilds dashboard summary tables incrementally, only for changed dates — including a "
            "full payment-mode breakdown (COD, prepaid, partially-prepaid) for both order counts "
            "and sales amounts.",
            "Tracks any checkout/session event type per product variant, not just checkout-initiated "
            "and buy-now — new event types are discovered automatically and get their own column, "
            "via a separate daily worker.",
            "Rolls up “bundle” product sales (TMC and PTS today) — bundle products are "
            "discovered automatically from the Shopify product catalog (active products only), "
            "not a hardcoded list.",
            "Sends live alerts to QStash: sales/session numbers and page-speed averages, per brand.",
            "Triggers the downstream cache pipeline directly after each run finishes, so "
            "dashboards pick up fresh data.",
        ],
    )
    add_body(
        doc,
        "Each brand is fully isolated — its own database, its own Shopify store, its own "
        "failure boundary. If one brand's run fails, the others still complete.",
    )

    add_h1(doc, "Architecture")
    add_image(doc, "architecture.png")

    add_h1(doc, "How it runs")
    add_bullets(
        doc,
        [
            "Automatically, every 10 minutes, via an internal scheduler (APScheduler).",
            "On demand, via POST /trigger (manual runs, or an external pipeline manager).",
            "Packaged as a Docker container, deployed on an EC2 instance.",
        ],
    )

    add_h1(doc, "Project structure")
    add_body(
        doc,
        "The pipeline used to be a single ~10,000-line file. It's now split into focused "
        "modules under pipeline/, all tied together through aws_background.py, which "
        "remains the entrypoint.",
    )
    add_image(doc, "project-structure.png")
    add_table(
        doc,
        ["Module", "Responsibility"],
        [
            ("aws_background.py", "Entrypoint, brand configuration, and the facade that re-exports everything below"),
            ("pipeline/state.py", "Shared config and long-lived objects (logger, Flask app, scheduler, brand registries)"),
            ("pipeline/db.py", "Database connections, retries, and cleanup"),
            ("pipeline/schema_migrations.py", "Table creation / schema upkeep"),
            ("pipeline/shopify_orders.py", "Order fetching, transforming, and loading"),
            ("pipeline/summary_incremental.py", "Rebuilding dashboard summary tables"),
            ("pipeline/mongo_events.py", "Checkout-initiated / buy-now event sync"),
            ("pipeline/web_vitals.py", "Page-speed data sync"),
            ("pipeline/shopifyql_sessions.py", "Session and city-level traffic analytics"),
            ("pipeline/utm_traffic.py", "UTM and referrer attribution"),
            ("pipeline/notifications.py", "QStash alerts and webhook triggers"),
            ("pipeline/orchestration.py", "Runs everything above, per brand, per pipeline tick"),
            ("bundle_rollups.py", "Bundle product discovery (Shopify Admin API, active products only) and sales rollups (TMC, PTS)"),
            ("workers/events_summary_worker.py", "Separate worker rolling up daily event totals, own Docker container"),
            ("workers/variant_wise_events_worker.py", "Separate worker rolling up any Mongo event type per product variant, own Docker container"),
            ("backfill scripts/", "Standalone scripts to re-run specific pieces over a historical date range"),
        ],
    )

    add_h1(doc, "Configuration")
    add_body(
        doc,
        "Brands are configured through numbered environment variables — one full set per "
        "brand index, e.g. BRAND_NAME_0, BRAND_TAG_0, DB_HOST_0, SHOP_NAME_0, ACCESS_TOKEN_0. "
        "In production, brand configuration is instead pulled from an internal brand-config API. "
        "Local/dev runs use TEST_MODE=true, which loads brand config from .env.local instead "
        "and disables outbound QStash calls.",
    )
    add_table(
        doc,
        ["Variable", "Purpose"],
        [
            ("SELF_RUN", "true = auto-run every 10 minutes; false = manual trigger only"),
            ("BACKFILL_MODE", "true = process a specific historical date window instead of “since last run”"),
            ("BACKFILL_BRAND_INDICES", "Restrict a backfill run to specific brands"),
            ("ALERTS_WEBHOOK_URL", "Where QStash-forwarded sales/web-vitals alerts get delivered"),
        ],
    )
    add_body(
        doc,
        "No secrets or real credentials are ever meant to live in source control — see "
        ".env.local / .env locally, and the deployment's own secret store in production.",
    )

    add_h1(doc, "Running it locally")
    add_body(doc, "pip install -r requirements.txt")
    add_body(doc, "python aws_background.py")
    add_body(
        doc,
        "With TEST_MODE=true, this loads brand config from .env.local, starts the scheduler, "
        "and serves the API on port 5000.",
    )
    add_body(doc, "Check it's alive:  curl http://localhost:5000/health")
    add_body(
        doc,
        "Trigger a run manually:  "
        "curl -X POST http://localhost:5000/trigger -H \"x-pipeline-key: <key>\" "
        "-d '{\"runId\": \"manual-test-1\"}'",
    )

    add_h1(doc, "API endpoints")
    add_table(
        doc,
        ["Endpoint", "Method", "Purpose"],
        [
            ("/health", "GET", "Basic liveness check"),
            ("/trigger", "POST", "Runs the pipeline immediately (requires x-pipeline-key + runId)"),
        ],
    )

    add_h1(doc, "Deployment")
    add_body(
        doc,
        "Runs as three Docker services (see docker-compose.yml): shopify-pipeline (the main "
        "pipeline), events-summary-worker (rolls up daily event totals across all brands, on "
        "its own daily schedule), and variant-wise-events-worker (rolls up any Mongo event "
        "type per product variant, into each brand's own database, on its own daily schedule).",
    )
    add_body(doc, "To deploy an update on the server:")
    add_bullets(
        doc,
        [
            "git pull origin master",
            "docker compose build shopify-pipeline",
            "docker compose up -d shopify-pipeline",
        ],
    )

    add_h2(doc, "Cache pipeline trigger")
    add_body(
        doc,
        "The cache pipeline runs as its own container (cache-pipeline) on the same EC2 host, "
        "joined to the same Docker network (pipeline-net) as shopify-pipeline. After each full "
        "run finishes, the pipeline fires a plain GET http://cache-pipeline:8080/trigger-pipeline "
        "— no QStash involved, just a direct container-to-container call over the shared network.",
    )
    add_body(
        doc,
        "This request runs in a background thread, so the pipeline doesn't wait for the cache "
        "rebuild to finish before moving on. Whenever the response does come back, only a short "
        "summary (status + brand count) is logged — not the full per-brand/per-date result body.",
    )

    add_h1(doc, "Testing")
    add_body(doc, "python -m unittest discover -p \"test_*.py\"")
    add_body(
        doc,
        "Covers brand-activation filtering, bundle rollup logic, UTM source mapping, "
        "and the events summary worker.",
    )

    add_h1(doc, "Backfill scripts")
    add_body(
        doc,
        "backfill scripts/ holds standalone tools for re-running one specific piece of the "
        "pipeline over a chosen date range — useful for fixing a gap in historical data "
        "without re-running everything. Each script is run directly with Python and reads "
        "the same brand configuration as the main pipeline.",
    )
    add_table(
        doc,
        ["Category", "Examples"],
        [
            ("Orders & addresses", "backfill_shopify.py, backfill_shopify_order_addresses.py, backfill_shopify_order_tags.py"),
            ("Returns / RTO", "backfill_returns_fact_order_created_date.py, backfill_rto_summary.py"),
            ("Checkout / buy-now events", "backfill_checkout_initiated_ci_events.py, backfill_buy_now_events.py (+ variant-level)"),
            ("Variant-wise events", "backfill_variant_wise_events_summary.py"),
            ("Sessions & web vitals", "backfill_hourly_sessions.py, backfill_web_vitals_summary.py, backfill_daily_web_vitals_summary.py"),
            ("UTM & traffic", "utm_sessions_backfill.py, referrer_sessions_backfill.py, backfill_product_traffic_split.py"),
            ("Bundles", "bundle_rollup_backfill.py"),
            ("Payment mode breakdown", "backfill_payment_mode_breakdown_columns.py — cod/prepaid/ppcod orders & sales only, on hour_wise_sales and overall_summary"),
        ],
    )

    doc.save(OUT_PATH)
    print("Wrote", OUT_PATH)


if __name__ == "__main__":
    build()
